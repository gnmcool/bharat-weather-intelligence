"""Government briefing exports for a state: Excel workbook and a print-ready HTML report
(use the browser's "Save as PDF"). Content = the same district table, official alerts and
methodology the app shows, so a Collector / SDMA briefing matches what is on screen."""
from __future__ import annotations

import html
import io
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

IST = ZoneInfo("Asia/Kolkata")
LEVEL = {0: "No risk", 1: "Watch", 2: "Alert", 3: "Severe"}
LEVEL_FILL = {0: "E8F5E9", 1: "FFF8E1", 2: "FFE0B2", 3: "FFCDD2"}
LEVEL_HEX = {0: "#2e7d32", 1: "#b58900", 2: "#e65100", 3: "#c62828"}

COLUMNS = [
    ("district", "District", 22), ("max_level", "Overall risk", 12), ("heat", "Heat", 9), ("cold", "Cold", 9),
    ("rain", "Heavy rain", 11), ("wind", "Wind", 9), ("tmax_7d_max", "Max temp 7d (°C)", 15),
    ("tmax_anom_7d", "Tmax vs normal (°C)", 17), ("tmin_7d_min", "Min temp 7d (°C)", 15),
    ("rain_7d", "Rain next 7d (mm)", 15), ("rain_7d_normal", "Normal 7d (mm)", 14), ("rain_max_day", "Max 1-day rain (mm)", 17),
    ("rain_30d", "Rain past 30d (mm)", 16), ("rain_30d_normal", "Normal 30d (mm)", 15), ("rain_30d_pct", "30d departure (%)", 15),
    ("rain_30d_cat", "30d IMD category", 16), ("gust_max", "Max gust (km/h)", 14), ("official_warnings", "Official alerts", 13),
    ("terrain", "Terrain class", 12),
]


def _row(d: dict[str, Any]) -> list[Any]:
    out = []
    for key, _, _ in COLUMNS:
        if key in ("heat", "cold", "rain", "wind"):
            out.append(LEVEL[d["levels"][key]])
        elif key == "max_level":
            out.append(LEVEL[d["max_level"]])
        else:
            out.append(d.get(key))
    return out


def _now() -> str:
    return datetime.now(IST).strftime("%d %b %Y, %I:%M %p IST")


def xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Districts"
    ws["A1"] = f"{data['state']} — district weather risk, next 7 days"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated {_now()} · Bharat Weather Intelligence · system-derived indicators (not official IMD warnings)"
    ws["A2"].font = Font(italic=True, size=9, color="666666")
    hdr = 4
    for c, (_, title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(hdr, c, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[hdr].height = 32
    for r, d in enumerate(data["districts"], hdr + 1):
        for c, v in enumerate(_row(d), 1):
            ws.cell(r, c, v)
        for c, key in ((2, "max_level"), (3, "heat"), (4, "cold"), (5, "rain"), (6, "wind")):
            lvl = d["max_level"] if key == "max_level" else d["levels"][key]
            ws.cell(r, c).fill = PatternFill("solid", fgColor=LEVEL_FILL[lvl])
    ws.freeze_panes = ws.cell(hdr + 1, 2)
    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(COLUMNS))}{hdr + len(data['districts'])}"

    wa = wb.create_sheet("Official alerts")
    heads = ["Severity", "Event", "Issuer", "Headline (verbatim)", "Area", "Effective", "Expires", "Urgency", "Certainty", "Link"]
    for c, h_ in enumerate(heads, 1):
        cell = wa.cell(1, c, h_)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7F1D1D")
    for r, w in enumerate(data["warnings"], 2):
        vals = [w.get("severity"), w.get("event"), w.get("issuer"), w.get("headline"), w.get("area"),
                w.get("effective"), w.get("expires"), w.get("urgency"), w.get("certainty"), w.get("link")]
        for c, v in enumerate(vals, 1):
            wa.cell(r, c, v)
    for c, wd in enumerate([10, 18, 16, 80, 30, 22, 22, 11, 11, 40], 1):
        wa.column_dimensions[get_column_letter(c)].width = wd
    for row in wa.iter_rows(min_row=2, max_col=4):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")

    wm = wb.create_sheet("Method & sources")
    lines = [
        ("Method", data["method"]),
        ("Levels", "0 No risk · 1 Watch · 2 Alert · 3 Severe — rules in docs/METHODOLOGY.md"),
        ("Heat", "IMD: Tmax ≥ 40 °C plains / 37 coastal / 30 hills and ≥ 4.5 °C above normal (≥ 6.5 severe); plains ≥ 45 / 47 °C"),
        ("Cold", "IMD: Tmin ≤ 10 °C plains / 15 coastal / 0 hills and ≤ −4.5 °C below normal; plains ≤ 4 / 2 °C"),
        ("Heavy rain", "IMD 24 h: ≥ 35.6 watch, ≥ 64.5 heavy (alert), ≥ 115.6 very heavy (severe)"),
        ("Wind", "Daily max gust: ≥ 50 watch, ≥ 62 alert, ≥ 89 km/h severe (Beaufort)"),
        ("30-day departure", "IMD classes vs 1991–2020 normal (NASA POWER); past 30 days = model analysis, not gauges"),
        ("Caveat", "System-derived indicators from forecast models. Not an IMD declaration. Official alerts are listed verbatim."),
    ] + [(f"Source {i + 1}", f"{s.get('source')} — {s.get('model') or ''} {s.get('notes') or ''}") for i, s in enumerate(data["sources"])]
    for r, (k, v) in enumerate(lines, 1):
        wm.cell(r, 1, k).font = Font(bold=True)
        wm.cell(r, 2, v).alignment = Alignment(wrap_text=True)
    wm.column_dimensions["A"].width = 18
    wm.column_dimensions["B"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def html_report(data: dict[str, Any]) -> str:
    e = html.escape
    ds = data["districts"]
    counts = {lvl: sum(1 for d in ds if d["max_level"] == lvl) for lvl in (3, 2, 1, 0)}
    warm = max((d for d in ds if d["tmax_anom_7d"] is not None), key=lambda d: d["tmax_anom_7d"], default=None)
    wet = max(ds, key=lambda d: d["rain_7d"], default=None)
    dry = min((d for d in ds if d["rain_30d_pct"] is not None), key=lambda d: d["rain_30d_pct"], default=None)
    concerns = [d for d in ds if d["max_level"] >= 2 or d["official_warnings"]]

    def cell(v: Any, nd: int = 1) -> str:
        if v is None:
            return "—"
        return f"{v:.{nd}f}" if isinstance(v, float) else e(str(v))

    def lvl(level: int) -> str:
        return f'<span class="pill" style="background:{LEVEL_HEX[level]}">{LEVEL[level]}</span>'

    rows = "".join(
        f"<tr><td class='l'>{e(d['district'])}</td><td>{lvl(d['max_level'])}</td>"
        f"<td>{cell(d['tmax_7d_max'])}</td><td>{cell(d['tmax_anom_7d'])}</td><td>{cell(d['tmin_7d_min'])}</td>"
        f"<td>{cell(d['rain_7d'], 0)}</td><td>{cell(d['rain_7d_normal'], 0)}</td><td>{cell(d['rain_max_day'], 0)}</td>"
        f"<td>{cell(d['rain_30d_pct'], 0)}</td><td>{cell(d['rain_30d_cat'])}</td><td>{cell(d['gust_max'], 0)}</td>"
        f"<td>{d['official_warnings'] or ''}</td></tr>"
        for d in ds
    )
    alerts = "".join(
        f"<li><b style='color:#b91c1c'>{e(w.get('severity') or '')}</b> · {e(w.get('event') or '')} · <i>{e(w.get('issuer') or '')}</i><br>"
        f"{e(w.get('headline') or '')}<br><small>Area: {e(w.get('area') or '—')} · valid to {e(str(w.get('expires') or '—'))}</small></li>"
        for w in data["warnings"]
    ) or "<li>No active official alerts matched to this state.</li>"
    focus = "".join(
        f"<li><b>{e(d['district'])}</b> — {LEVEL[d['max_level']]}"
        + "".join(f"; {k} {LEVEL[v].lower()}" for k, v in d["levels"].items() if v)
        + (f"; {d['official_warnings']} official alert(s)" if d["official_warnings"] else "") + "</li>"
        for d in concerns[:15]
    ) or "<li>No district at Alert or Severe level in the next 7 days.</li>"
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>{e(data['state'])} weather risk briefing</title>
<style>
 @page {{ size: A4 landscape; margin: 12mm; }}
 body {{ font: 12px/1.4 system-ui, Segoe UI, sans-serif; color:#111; margin: 18px; }}
 h1 {{ font-size: 20px; margin: 0; }} h2 {{ font-size: 14px; margin: 18px 0 6px; border-bottom: 2px solid #1f3a5f; padding-bottom: 2px; }}
 .sub {{ color:#555; font-size: 11px; }} .grid {{ display:flex; gap:10px; margin-top:10px; }}
 .box {{ flex:1; border:1px solid #ddd; border-radius:8px; padding:8px; text-align:center; }}
 .box b {{ font-size: 22px; display:block; }}
 table {{ width:100%; border-collapse: collapse; font-size: 11px; }} th {{ background:#1f3a5f; color:#fff; padding:4px; }}
 td {{ border-bottom:1px solid #eee; padding:3px 4px; text-align:center; }} td.l {{ text-align:left; font-weight:600; }}
 tr:nth-child(even) td {{ background:#fafafa; }}
 .pill {{ color:#fff; border-radius:8px; padding:1px 6px; font-size:10px; }}
 .note {{ font-size:10px; color:#555; }} ul {{ margin: 4px 0; padding-left: 18px; }} li {{ margin-bottom: 4px; }}
 .bar {{ display:flex; justify-content:space-between; align-items:flex-start; }}
 .print {{ background:#1f3a5f; color:#fff; border:0; border-radius:6px; padding:8px 14px; cursor:pointer; }}
 @media print {{ .print {{ display:none; }} body {{ margin:0; }} }}
</style></head><body>
<div class="bar"><div><h1>{e(data['state'])} — Weather risk briefing (next 7 days)</h1>
<div class="sub">Generated {_now()} · Bharat Weather Intelligence · Made by Gaurav Makwana</div></div>
<button class="print" onclick="window.print()">Print / Save as PDF</button></div>
<div class="grid">
 <div class="box" style="border-color:{LEVEL_HEX[3]}"><b style="color:{LEVEL_HEX[3]}">{counts[3]}</b>Severe</div>
 <div class="box" style="border-color:{LEVEL_HEX[2]}"><b style="color:{LEVEL_HEX[2]}">{counts[2]}</b>Alert</div>
 <div class="box" style="border-color:{LEVEL_HEX[1]}"><b style="color:{LEVEL_HEX[1]}">{counts[1]}</b>Watch</div>
 <div class="box"><b>{len(data['warnings'])}</b>Official alerts</div>
</div>
<h2>Where is something unusual?</h2><ul>
{f"<li>Warmest vs normal: <b>{e(warm['district'])}</b> {warm['tmax_anom_7d']:+.1f} °C (7-day mean max temperature)</li>" if warm else ""}
{f"<li>Wettest next 7 days: <b>{e(wet['district'])}</b> {wet['rain_7d']:.0f} mm" + (f" (normal {wet['rain_7d_normal']:.0f} mm)" if wet.get('rain_7d_normal') is not None else "") + "</li>" if wet else ""}
{f"<li>Largest 30-day rainfall deficit: <b>{e(dry['district'])}</b> {dry['rain_30d_pct']:+.0f}% ({e(dry['rain_30d_cat'] or '')})</li>" if dry else ""}
</ul>
<h2>Districts needing attention</h2><ul>{focus}</ul>
<h2>Official alerts (verbatim, NDMA SACHET — IMD / CWC / SDMA)</h2><ul>{alerts}</ul>
<h2>All districts</h2>
<table><thead><tr><th>District</th><th>Risk</th><th>Max °C</th><th>Δ vs normal</th><th>Min °C</th><th>Rain 7d mm</th><th>Normal 7d</th><th>Max 1-day</th><th>30d %</th><th>30d class</th><th>Gust km/h</th><th>Alerts</th></tr></thead>
<tbody>{rows}</tbody></table>
<h2>Method & limits</h2>
<p class="note">{e(data['method'])} Risk levels are system-derived indicators computed from forecast models (ECMWF/GFS via Open-Meteo) using IMD criteria;
they are not IMD declarations. Normals: NASA POWER 1991–2020. Past 30-day rain is model analysis, not rain-gauge data.
Official alerts above are reproduced verbatim from the NDMA SACHET CAP feed. Always act on official IMD / SDMA instructions.</p>
</body></html>"""
